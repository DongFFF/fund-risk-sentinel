"""``.docx`` 材料解析（python-docx）。

报销单、审批意见、合同等业务材料大量以 Word 形式流转，因此这里是使用频率
最高的 loader。除了正文段落，还抽取表格、页眉页脚与文档属性——审批意见常常
藏在页眉的签批栏或表格的最后一行。

注意：``.doc``（旧二进制格式）python-docx 无法读取，会给出明确的转换提示。
"""

from __future__ import annotations

from pathlib import Path

from ...core.enums import DocumentFormat, DocumentRole
from ...core.exceptions import CorruptedFileError, IngestError
from ...core.models import Material
from ..base import BaseLoader


class DocxLoader(BaseLoader):
    """Word 文档解析器。"""

    formats = (DocumentFormat.DOCX,)
    extensions = (".docx", ".doc")
    requires = ("python-docx",)

    def load(self, path: Path | str, *, role: DocumentRole = DocumentRole.OTHER) -> Material:
        p = Path(path)
        if p.suffix.lower() == ".doc":
            raise IngestError(
                f"{p.name} 是旧版 .doc 格式，无法直接解析。"
                "请用 Word / WPS 另存为 .docx 后重新上传。"
            )

        try:
            import docx  # noqa: PLC0415 - 延迟导入，避免未装依赖时影响其他格式
            from docx.document import Document as _DocxDocument
            from docx.oxml.table import CT_Tbl
            from docx.oxml.text.paragraph import CT_P
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise self._missing_dep("python-docx", "Word 文档") from exc

        try:
            document: "_DocxDocument" = docx.Document(str(p))
        except Exception as exc:  # noqa: BLE001 - python-docx 抛的异常类型不稳定
            raise self._wrap_error(p, exc) from exc

        material = self._base_material(p, role, DocumentFormat.DOCX)
        blocks: list[str] = []
        tables: list[list[list[str]]] = []
        heading_outline: list[dict[str, str]] = []

        # 按文档流顺序遍历段落与表格，保持原始阅读顺序
        for block in _iter_blocks(document):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                style = (block.style.name if block.style is not None else "") or ""
                if style.lower().startswith("heading") or style.startswith("标题"):
                    heading_outline.append({"style": style, "text": text})
                    blocks.append(f"【{text}】")
                else:
                    blocks.append(text)
            elif isinstance(block, Table):
                rows = _table_to_rows(block)
                if rows:
                    tables.append(rows)
                    blocks.append(_render_docx_table(rows, len(tables)))

        for section_index, section in enumerate(document.sections, start=1):
            for label, container in (("页眉", section.header), ("页脚", section.footer)):
                texts = [para.text.strip() for para in container.paragraphs if para.text.strip()]
                if texts:
                    blocks.append(f"【{label}{section_index}】" + " ".join(texts))

        material.text = "\n".join(blocks)
        material.tables = tables
        material.metadata["paragraph_count"] = len(document.paragraphs)
        material.metadata["table_count"] = len(tables)
        material.metadata["headings"] = heading_outline
        material.metadata["sections"] = len(document.sections)

        try:
            props = document.core_properties
            material.metadata["core_properties"] = {
                "author": props.author or "",
                "last_modified_by": props.last_modified_by or "",
                "created": props.created.isoformat() if props.created else "",
                "modified": props.modified.isoformat() if props.modified else "",
                "title": props.title or "",
            }
        except Exception:  # noqa: BLE001 - 属性缺失不影响正文解析
            pass

        # 用分页符数量估算页数
        material.page_count = max(1, sum(1 for para in document.paragraphs if "w:br" in para._p.xml and 'w:type="page"' in para._p.xml) + 1)
        return self._finalize(material)


def _iter_blocks(document):
    """按文档流顺序产出段落与表格。"""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _table_to_rows(table) -> list[list[str]]:
    """表格 → 二维字符串数组，处理合并单元格的重复取值。"""
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            # 合并单元格会被重复引用，去重避免同一内容重复出现
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            cells.append(cell.text.strip().replace("\n", " "))
        if any(cells):
            rows.append(cells)
    return rows


def _render_docx_table(rows: list[list[str]], index: int) -> str:
    """把表格渲染为便于模型阅读与定位的文本。"""
    lines = [f"【表格{index}】"]
    header = rows[0]
    for i, row in enumerate(rows[1:], start=1):
        pairs = [
            f"{header[j] if j < len(header) else f'列{j + 1}'}：{cell}"
            for j, cell in enumerate(row)
        ]
        lines.append(f"第{i}行 | " + "；".join(pairs))
    if len(rows) == 1:
        lines.append(" | ".join(header))
    lines.append(f"（表格{index}共 {max(0, len(rows) - 1)} 行）")
    return "\n".join(lines)


__all__ = ["DocxLoader"]
