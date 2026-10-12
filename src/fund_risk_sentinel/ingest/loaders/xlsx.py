"""``.xlsx`` 材料解析（openpyxl）。

经费台账、预算执行表、报销明细这类材料本质是表格，风险线索往往藏在
"行与行之间"（同一供应商多笔、金额累加超预算）。因此这里保留完整的二维表结构，
而不是只取一段文字。
"""

from __future__ import annotations

from pathlib import Path

from ...core.enums import DocumentFormat, DocumentRole
from ...core.exceptions import CorruptedFileError, IngestError
from ...core.models import Material
from ..base import BaseLoader

_MAX_SHEETS = 20


class XlsxLoader(BaseLoader):
    """Excel 工作簿解析器。"""

    formats = (DocumentFormat.XLSX,)
    extensions = (".xlsx", ".xlsm", ".xls")
    requires = ("openpyxl",)

    def load(self, path: Path | str, *, role: DocumentRole = DocumentRole.OTHER) -> Material:
        p = Path(path)
        if p.suffix.lower() == ".xls":
            raise IngestError(
                f"{p.name} 是旧版 .xls 格式，openpyxl 无法读取。"
                "请用 Excel / WPS 另存为 .xlsx 后重新上传。"
            )
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise self._missing_dep("openpyxl", "Excel 表格") from exc

        try:
            workbook = load_workbook(
                filename=str(p), read_only=True, data_only=True, keep_links=False
            )
        except Exception as exc:  # noqa: BLE001
            raise self._wrap_error(p, exc) from exc

        material = self._base_material(p, role, DocumentFormat.XLSX)
        blocks: list[str] = []
        tables: list[list[list[str]]] = []
        sheet_meta: list[dict[str, object]] = []

        try:
            sheet_names = workbook.sheetnames[:_MAX_SHEETS]
            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                rows: list[list[str]] = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index >= self.context.max_rows:
                        rows.append(["……", f"（超过 {self.context.max_rows} 行，已截断）"])
                        break
                    values = ["" if cell is None else str(cell).strip() for cell in row]
                    if any(values):
                        rows.append(values)
                if not rows:
                    continue
                tables.append(rows)
                blocks.append(_render_sheet(rows, sheet_name))
                sheet_meta.append({"name": sheet_name, "rows": len(rows), "cols": len(rows[0])})
        finally:
            workbook.close()

        material.text = "\n\n".join(blocks)
        material.tables = tables
        material.metadata["sheets"] = sheet_meta
        material.metadata["sheet_count"] = len(sheet_meta)
        if len(workbook.sheetnames) > _MAX_SHEETS:
            material.metadata["sheets_truncated"] = True
        material.page_count = max(1, len(sheet_meta))
        return self._finalize(material)


def _render_sheet(rows: list[list[str]], sheet_name: str) -> str:
    """工作表 → 文本。首行视作表头，其余按「键：值」逐行展开。"""
    lines = [f"【工作表：{sheet_name}】"]
    header = rows[0]
    body = rows[1:]
    if not body:
        lines.append(" | ".join(header))
        return "\n".join(lines)

    # 判断首行是否像表头：非空且不与数据行同型（不全是数字）
    looks_like_header = any(
        not _is_number(cell) for cell in header if cell
    ) and sum(1 for c in header if c) >= max(2, len(header) // 2)

    if looks_like_header:
        for i, row in enumerate(body, start=1):
            pairs = [
                f"{header[j] if j < len(header) else f'列{j + 1}'}：{cell}"
                for j, cell in enumerate(row)
                if cell
            ]
            lines.append(f"第{i}行 | " + "；".join(pairs))
        lines.append(f"（{sheet_name} 共 {len(body)} 条记录）")
    else:
        for i, row in enumerate(rows, start=1):
            lines.append(f"第{i}行 | " + " | ".join(row))
    return "\n".join(lines)


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", ""))
        return True
    except ValueError:
        return False


__all__ = ["XlsxLoader"]
